"""PDF to Excel converter using pdfplumber and openpyxl."""

import argparse
import sys
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def extract_tables_from_pdf(pdf_path: str) -> list[dict]:
    """PDFから全ページのテーブルを抽出する。

    Returns:
        各ページのテーブルデータを含むリスト。
        [{"page": int, "tables": [[[str, ...], ...], ...]}, ...]
    """
    results = []
    with pdfplumber.open(pdf_path) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables()
            if tables:
                results.append({"page": page_num, "tables": tables})
            else:
                # テーブルが見つからない場合、テキストを行ごとに抽出
                text = page.extract_text()
                if text:
                    lines = [[line] for line in text.split("\n") if line.strip()]
                    if lines:
                        results.append({"page": page_num, "tables": [lines]})
    return results


def write_to_excel(data: list[dict], output_path: str, separate_sheets: bool = False):
    """抽出したデータをExcelファイルに書き込む。

    Args:
        data: extract_tables_from_pdf の戻り値。
        output_path: 出力先 Excel ファイルパス。
        separate_sheets: True ならページごとに別シートに書き込む。
    """
    wb = Workbook()
    wb.remove(wb.active)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    if not data:
        ws = wb.create_sheet("Sheet1")
        ws["A1"] = "テーブルが見つかりませんでした"
        wb.save(output_path)
        return

    if separate_sheets:
        for page_data in data:
            page_num = page_data["page"]
            for table_idx, table in enumerate(page_data["tables"]):
                sheet_name = f"Page{page_num}"
                if table_idx > 0:
                    sheet_name = f"Page{page_num}_Table{table_idx + 1}"
                # シート名は31文字まで
                ws = wb.create_sheet(sheet_name[:31])
                _write_table(ws, table, header_font, header_fill, header_alignment, thin_border)
    else:
        ws = wb.create_sheet("Sheet1")
        current_row = 1
        for page_data in data:
            page_num = page_data["page"]
            for table_idx, table in enumerate(page_data["tables"]):
                # ページ区切りヘッダー
                if current_row > 1:
                    current_row += 1
                cell = ws.cell(row=current_row, column=1, value=f"--- Page {page_num} ---")
                cell.font = Font(bold=True, size=12)
                current_row += 1

                current_row = _write_table(
                    ws, table, header_font, header_fill, header_alignment, thin_border,
                    start_row=current_row,
                )

    # 列幅の自動調整
    for ws in wb.worksheets:
        for col in ws.columns:
            max_length = 0
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            adjusted_width = min(max_length + 4, 50)
            ws.column_dimensions[col[0].column_letter].width = max(adjusted_width, 8)

    wb.save(output_path)


def _write_table(ws, table, header_font, header_fill, header_alignment, thin_border,
                 start_row=1) -> int:
    """ワークシートにテーブルデータを書き込む。書き込み後の次の行番号を返す。"""
    for row_idx, row in enumerate(table):
        for col_idx, value in enumerate(row):
            cell = ws.cell(
                row=start_row + row_idx,
                column=col_idx + 1,
                value=value if value else "",
            )
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

            # 最初の行をヘッダーとしてスタイル
            if row_idx == 0:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment

    return start_row + len(table)


def convert(input_path: str, output_path: str | None = None, separate_sheets: bool = False):
    """PDFファイルをExcelに変換するメイン関数。

    Args:
        input_path: 入力PDFファイルパス。
        output_path: 出力Excelファイルパス（省略時は入力ファイルと同名の.xlsx）。
        separate_sheets: ページごとに別シートにするか。
    """
    input_file = Path(input_path)
    if not input_file.exists():
        raise FileNotFoundError(f"ファイルが見つかりません: {input_path}")
    if not input_file.suffix.lower() == ".pdf":
        raise ValueError(f"PDFファイルを指定してください: {input_path}")

    if output_path is None:
        output_path = str(input_file.with_suffix(".xlsx"))

    print(f"変換中: {input_path} -> {output_path}")
    data = extract_tables_from_pdf(input_path)
    write_to_excel(data, output_path, separate_sheets)

    total_tables = sum(len(p["tables"]) for p in data)
    total_pages = len(data)
    print(f"完了: {total_pages}ページから{total_tables}個のテーブルを抽出しました")
    print(f"出力: {output_path}")


def main():
    parser = argparse.ArgumentParser(description="PDFファイルをExcel(.xlsx)に変換します")
    parser.add_argument("input", help="入力PDFファイルのパス")
    parser.add_argument("-o", "--output", help="出力Excelファイルのパス（デフォルト: 入力ファイルと同名の.xlsx）")
    parser.add_argument(
        "-s", "--separate-sheets",
        action="store_true",
        help="ページごとに別シートに出力する",
    )
    args = parser.parse_args()

    try:
        convert(args.input, args.output, args.separate_sheets)
    except (FileNotFoundError, ValueError) as e:
        print(f"エラー: {e}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
